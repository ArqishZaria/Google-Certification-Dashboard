"""Django integration tests: scope enforcement, import round-trip, views.  Run: python manage.py test"""
import csv
import os
import tempfile
from datetime import datetime

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from . import services
from .logic.tests_logic import master_row, prog_row, spec_row
from .models import Participant, Region, UserScope


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def make_files(tmp):
    from openpyxl import Workbook
    m = [master_row(1, "4210100000001", "a@x.com", region="Southern Region"),
         master_row(2, "7140100000002", "b@x.com", region="Hunza Region", **{"Current Local Council": "Altit & Karimabad", "JK Attendance": "ALTIT"})]
    wb = Workbook(); ws = wb.active; ws.title = "Certificate Applications"
    cols = list(m[0].keys()); ws.append(cols)
    for r in m:
        ws.append([r[c] for c in cols])
    wb.save(os.path.join(tmp, "master.xlsx"))
    wb = Workbook(); ws = wb.active; ws.title = "Southern"
    ws.append(["id", "License Issued Category", "Full Name", "Email Address (Course Access)", "CNIC Number", "Gender"])
    ws.append([1, "Southern", "Person", "a@x.com", "4210100000001", "Male"])
    ws2 = wb.create_sheet("Hunza")
    ws2.append(["id", "License Issued Category", "Full Name", "Email Address (Course Access)", "CNIC Number", "Gender"])
    ws2.append([2, "Hunza Region", "Person B", "b@x.com", "7140100000002", "Male"])
    wb.save(os.path.join(tmp, "final.xlsx"))
    acct = "gwg-ent-testacct"
    suffix = f"{acct}_-_CourseraEnterpriseExport_2026-09-16_05-42-47_UTC.csv"
    write_csv(os.path.join(tmp, f"Program_Activity_{suffix}"), [
        prog_row("a@x.com", "4210100000001"), prog_row("b@x.com", None, "INVITED", "AKEPB - Hunza (Male)")])
    write_csv(os.path.join(tmp, f"Specialisation_Activity_{suffix}"), [spec_row("a@x.com", done="9", total="9", completed="Yes")])
    write_csv(os.path.join(tmp, f"Course_Activity_{suffix}"), [{
        "Name": "Person", "Email": "a@x.com", "External ID": "4210100000001", "Course": "Foundations", "Course ID": "c1", "Course Slug": "f",
        "Enrollment Time": "2026-08-26T00:00:00Z", "Last Course Activity Time": "2026-09-10T00:00:00Z", "Overall Progress": "100",
        "Completed": "Yes", "Removed From Program": "No", "Completion Time": "2026-09-10T00:00:00Z", "Learning Hours": "3", "Course Type": "Course"}])
    return sorted(os.path.join(tmp, f) for f in os.listdir(tmp))


class ImportAndAccessTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.paths = make_files(cls.tmp)
        cls.imp = services.run_import(cls.paths)

    def test_import_creates_linked_records(self):
        self.assertEqual(Participant.objects.count(), 2)
        a = Participant.objects.get(email="a@x.com")
        self.assertEqual((a.status, a.pc_progress, a.region.name), ("COMPLETED", 100.0, "Southern"))
        self.assertIsNotNone(a.registration)
        self.assertEqual(a.licenses.count(), 1)
        b = Participant.objects.get(email="b@x.com")
        self.assertEqual((b.status, b.license_state, b.region.name), ("PENDING", "PENDING", "Hunza"))

    def test_import_is_idempotent(self):
        services.run_import(self.paths)
        self.assertEqual(Participant.objects.count(), 2)

    def test_recompute_matches_import(self):
        before = {p.key: (p.status, p.pc_progress) for p in Participant.objects.all()}
        services.recompute_all()
        self.assertEqual(before, {p.key: (p.status, p.pc_progress) for p in Participant.objects.all()})

    def test_changing_category_changes_completion(self):
        from .models import Specialization
        Specialization.objects.filter(slug="google-cybersecurity").update(category="SHORT")
        services.recompute_all()
        self.assertEqual(Participant.objects.get(email="a@x.com").status, "NO_PC")

    def test_scope_regional_user_sees_only_own_region(self):
        u = User.objects.create_user("southchair", password="pw")
        UserScope.objects.create(user=u, role="REGIONAL_CHAIR", region=Region.objects.get(name="Southern"))
        self.client.login(username="southchair", password="pw")
        r = self.client.get(reverse("participants"))
        self.assertContains(r, "Person")
        self.assertNotContains(r, "Person B")
        # trying to widen the scope through the URL must not leak other regions
        r = self.client.get(reverse("participants") + "?region=Hunza")
        self.assertNotContains(r, "Person B")
        hunza = Participant.objects.get(email="b@x.com")
        self.assertEqual(self.client.get(reverse("participant_detail", args=[hunza.public_id])).status_code, 404)
        self.assertEqual(self.client.get(reverse("data_quality")).status_code, 403)

    def test_programme_filter_and_excel_view_render(self):
        User.objects.create_superuser("root2", "r2@x.com", "pw")
        self.client.login(username="root2", password="pw")
        r = self.client.get(reverse("dashboard") + "?programme=regional")
        self.assertContains(r, "Excel-compatible view")
        self.assertEqual(self.client.get(reverse("dashboard") + "?programme=national").status_code, 200)
        self.assertEqual(self.client.get(reverse("report_download", args=["excel-compatible", "xlsx"])).status_code, 200)

    def test_region_is_residence_based(self):
        self.assertEqual(Participant.objects.get(email="b@x.com").region.name, "Hunza")

    def test_user_without_scope_gets_403(self):
        User.objects.create_user("nobody", password="pw")
        self.client.login(username="nobody", password="pw")
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 403)
        self.assertEqual(self.client.get(reverse("report_download", args=["participant", "csv"])).status_code, 403)

    def test_anonymous_redirected_to_login(self):
        r = self.client.get(reverse("dashboard"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("login", r["Location"])

    def test_superuser_sees_everything_and_downloads_work(self):
        User.objects.create_superuser("root", "r@x.com", "pw")
        self.client.login(username="root", password="pw")
        self.assertContains(self.client.get(reverse("dashboard")), "National Dashboard")
        for slug in ("national", "participant", "completion", "followup-pending"):
            for fmt in ("csv", "xlsx", "pdf"):
                self.assertEqual(self.client.get(reverse("report_download", args=[slug, fmt])).status_code, 200, (slug, fmt))
        self.assertEqual(self.client.get(reverse("followup") + "?list=pending").status_code, 200)
        self.assertEqual(self.client.get(reverse("data_quality")).status_code, 200)
        self.assertEqual(self.client.get(reverse("methodology")).status_code, 200)
        self.assertEqual(self.client.get(reverse("participants") + "?q=person&date_type=invited&date_from=garbage").status_code, 200)
