"""Prove the engine reproduces the 2025 Excel dashboard.

Feeds the OLD workbook's `Data` sheet through the same derive/metrics code that powers the new dashboard and
compares the result with the values printed on the workbook's `National Dashboard` sheet (cached pivot results).
Run:  python manage.py verify_excel_parity "GCC Dashboard as of 5th Nov 2025 - National.xlsx"
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date

from . import catalog as C
from . import metrics as M
from . import normalize as N
from . import readers as R

# Excel cell -> meaning (National Dashboard sheet)
EXCEL_CELLS = {
    "D7": "Candidates invited", "H7": "Candidates started any one course", "L7": "No. of courses registered",
    "U7": "Candidates not started any course", "M29": "Candidates completed 1 course", "P29": "Total no. of courses completed",
    "C52": "Candidates >50% completion", "F52": "Courses >50% completion", "L52": "Candidates <50% completion",
    "O52": "Courses <50% completion", "V52": "Candidates on 0% completion", "Y52": "Courses on 0% completion",
    "E8": "Invited - Male", "E9": "Invited - Female", "I8": "Started - Male", "I9": "Started - Female",
    "V8": "Not started - Male", "V9": "Not started - Female",
}


def compute_from_old_data(data_rows, as_of=date(2025, 11, 5), band_rule="best"):
    """band_rule: "best" = most advanced specialization (dashboard default);
    "last_row" = reproduce the Excel artifact (its per-candidate flag sits on the candidate's LAST row in the sheet)."""
    by_person = defaultdict(lambda: {"specs": [], "gender": None})
    for r in data_rows:
        email = N.norm_email(r.get("Email"))
        if not email:
            continue
        p = by_person[email]
        g = N.title_gender(r.get("Gender"))
        p["gender"] = p["gender"] or g
        if N.clean_text(r.get("Specialization")):
            done = N.is_yes(r.get("Completed"))
            cc, tc = N.parse_int(r.get("# Completed Courses"), 0), N.parse_int(r.get("# Courses in Specialization"), 0)
            fam, cat, _ = C.classify_specialization(N.clean_text(r.get("Specialization Slug")) or "", N.clean_text(r.get("Specialization")))
            p["specs"].append({"family": fam, "category": cat, "completed": done, "progress_pct": 100.0 if done else (cc * 100.0 / tc if tc else 0.0),
                               "completed_at": None, "enrolled_at": None, "last_activity_at": None})
    parts, spec_rows = [], []
    for email, p in by_person.items():
        d = M.derive_participant([{"state": "ACCEPTED", "removed": False}], p["specs"], [], as_of)
        if band_rule == "last_row" and p["specs"]:
            last = p["specs"][-1]
            d["legacy_band"] = C.legacy_band(100.0 if last["completed"] else last["progress_pct"], last["completed"])
        d.update({"key": email, "gender": p["gender"], "shortlisted": True})
        parts.append(d)
        spec_rows += [{"participant_key": email, **s} for s in p["specs"]]
    return M.legacy_kpis(parts, spec_rows)


def verify(path):
    """-> rows of (label, cell, excel value, engine value [default rule], ok, engine value [Excel-exact rule], ok_exact)"""
    from openpyxl import load_workbook
    data = R.read_xlsx_sheets(path)["Data"]
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["National Dashboard"]
    excel = {cell: ws[cell].value for cell in EXCEL_CELLS}
    wb.close()

    def mapping(o):
        g = o["gender"]
        return {
            "D7": o["invited"], "H7": o["started"], "L7": o["courses_registered"], "U7": o["not_started"],
            "M29": o["completed_candidates"], "P29": o["courses_completed"],
            "C52": o["gt50_candidates"], "F52": o["gt50_courses"], "L52": o["lt50_candidates"], "O52": o["lt50_courses"],
            "V52": o["zero_candidates"], "Y52": o["zero_courses"],
            "E8": g["Male"]["invited"], "E9": g["Female"]["invited"], "I8": g["Male"]["started"], "I9": g["Female"]["started"],
            "V8": g["Male"]["not_started"], "V9": g["Female"]["not_started"],
        }
    best = mapping(compute_from_old_data(data, band_rule="best"))
    exact = mapping(compute_from_old_data(data, band_rule="last_row"))
    return [(EXCEL_CELLS[c], c, excel[c], best[c], excel[c] == best[c], exact[c], excel[c] == exact[c]) for c in EXCEL_CELLS]


def verify_regions(path):
    """Per-region cards of the 'National Pivot' sheet vs the engine (Excel-exact band rule).
    Region = the form's residing regional council, exactly like the workbook."""
    from openpyxl import load_workbook
    data = R.read_xlsx_sheets(path)["Data"]
    by_person = defaultdict(lambda: {"specs": [], "gender": None, "region": None})
    for r in data:
        email = N.norm_email(r.get("Email"))
        if not email:
            continue
        p = by_person[email]
        p["region"] = p["region"] or C.canon_region(r.get("Currently residing in which Regional Council?"))
        if N.clean_text(r.get("Specialization")):
            done = N.is_yes(r.get("Completed"))
            cc, tc = N.parse_int(r.get("# Completed Courses"), 0), N.parse_int(r.get("# Courses in Specialization"), 0)
            p["specs"].append({"family": N.clean_text(r.get("Specialization")), "category": C.PC, "completed": done,
                               "progress_pct": 100.0 if done else (cc * 100.0 / tc if tc else 0.0),
                               "completed_at": None, "enrolled_at": None, "last_activity_at": None})
    parts, rows = [], []
    for email, p in by_person.items():
        d = M.derive_participant([{"state": "ACCEPTED", "removed": False}], p["specs"], [], date(2025, 11, 5))
        if p["specs"]:
            last = p["specs"][-1]
            d["legacy_band"] = C.legacy_band(100.0 if last["completed"] else last["progress_pct"], last["completed"])
        d.update({"key": email, "gender": None, "shortlisted": True, "region": p["region"] or "Unassigned"})
        parts.append(d)
        rows += [{"participant_key": email, **s} for s in p["specs"]]
    table = {r["name"]: r for r in M.legacy_group_table(parts, rows, "region")}

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["National Pivot"]
    short = {"Ishkoman Puniyal": "Ishkoman Puniyal"}

    def block(col_label, first_row, last_row, cols):
        out = {}
        for r in range(first_row, last_row + 1):
            label = ws[f"{col_label}{r}"].value
            if label and label != "Grand Total" and label != "(blank)":
                out[label] = [ws[f"{c}{r}"].value for c in cols]
        return out
    inv = block("A", 4, 12, ["B", "C", "D"])            # invited, started, courses registered
    comp = block("A", 34, 42, ["B", "C"])                # candidates / courses completed
    lt50 = block("G", 34, 42, ["H", "I"])               # National Dashboard L52/O52 read from $G$33
    zero = block("K", 34, 42, ["L", "M"])               # V52/Y52 read from $K$33
    gt50 = block("O", 34, 42, ["P", "Q"])               # C52/F52 read from $O$33
    wb.close()
    out = []
    for region in inv:
        t = table.get(region, {})
        checks = [("invited", inv[region][0], t.get("invited", 0)), ("started", inv[region][1], t.get("started", 0)),
                  ("courses registered", inv[region][2], t.get("courses_registered", 0)),
                  ("cand. 100%", comp.get(region, [0, 0])[0], t.get("completed_candidates", 0)), ("courses 100%", comp.get(region, [0, 0])[1], t.get("courses_completed", 0)),
                  ("cand. >50%", gt50.get(region, [0, 0])[0], t.get("gt50_candidates", 0)), ("courses >50%", gt50.get(region, [0, 0])[1], t.get("gt50_courses", 0)),
                  ("cand. <50%", lt50.get(region, [0, 0])[0], t.get("lt50_candidates", 0)), ("courses <50%", lt50.get(region, [0, 0])[1], t.get("lt50_courses", 0)),
                  ("cand. 0%", zero.get(region, [0, 0])[0], t.get("zero_candidates", 0)), ("courses 0%", zero.get(region, [0, 0])[1], t.get("zero_courses", 0))]
        out += [(region, label, e, g, (e or 0) == g) for label, e, g in checks]
    return out
