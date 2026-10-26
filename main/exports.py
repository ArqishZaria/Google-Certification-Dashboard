"""CSV / Excel / PDF exporters. Cells that could be interpreted as spreadsheet formulas are neutralised."""
import csv
import io
from datetime import date
from xml.sax.saxutils import escape as _esc

from django.http import HttpResponse

from .logic.sanitize import cell_value as _cell, sheet_title


def _filename(title, ext):
    safe = "".join(ch if ch.isalnum() else "_" for ch in title).strip("_")
    return f"GCC_{safe}_{date.today():%Y%m%d}.{ext}"


def to_csv(title, columns, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    for r in rows:
        w.writerow([_cell(v) for v in r])
    resp = HttpResponse("\ufeff" + buf.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(title, "csv")}"'
    return resp


def to_xlsx(title, subtitle, columns, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title(title)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="666666")
    hdr = 4
    for j, c in enumerate(columns, 1):
        cell = ws.cell(row=hdr, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, hdr + 1):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=_cell(v))
    for j, c in enumerate(columns, 1):
        width = max([len(str(c))] + [len(str(_cell(r[j - 1]))) for r in rows[:300]] + [8])
        ws.column_dimensions[get_column_letter(j)].width = min(width + 2, 45)
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(columns))}{hdr + max(len(rows), 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(title, "xlsx")}"'
    return resp


def to_pdf(title, subtitle, columns, rows):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm, title=title)
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small")
    small.fontSize = 6.5 if len(columns) > 9 else 7.5
    small.leading = small.fontSize + 2
    head = styles["BodyText"].clone("head")
    head.fontSize = small.fontSize
    head.textColor = colors.white
    head.leading = small.leading

    def p(v, st):
        return Paragraph(_esc(str(_cell(v))), st)

    data = [[p(c, head) for c in columns]] + [[p(v, small) for v in r] for r in rows]
    width = landscape(A4)[0] - 20 * mm
    t = Table(data, colWidths=[width / len(columns)] * len(columns), repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D1DC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    doc.build([Paragraph(_esc(title), styles["Title"]), Paragraph(_esc(subtitle), styles["Italic"]), Spacer(1, 4 * mm), t])
    resp = HttpResponse(buf.getvalue(), content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(title, "pdf")}"'
    return resp


def respond(fmt, title, subtitle, columns, rows):
    if fmt == "csv":
        return to_csv(title, columns, rows)
    if fmt == "xlsx":
        return to_xlsx(title, subtitle, columns, rows)
    if fmt == "pdf":
        return to_pdf(title, subtitle, columns, rows)
    raise ValueError(fmt)
